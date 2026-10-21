#!/usr/bin/env python3
"""Az Excel alaprajzból (raktar_alaprajz_anonim.xlsx) előállítja a data/adatok.js fájlt.

Használat (a repó gyökeréből):
    pip install openpyxl
    python3 tools/xlsx_to_adatok.py [bemenet.xlsx] [kimenet.js]

A script a cellák *mentett értékeit* olvassa, ezért képletes munkafüzetnél az
Excelben mentés után futtasd (a mentés frissíti a gyorsítótárazott értékeket).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

GYOKER = Path(__file__).resolve().parent.parent
BEMENET = Path(sys.argv[1]) if len(sys.argv) > 1 else GYOKER / "raktar_alaprajz_anonim.xlsx"
KIMENET = Path(sys.argv[2]) if len(sys.argv) > 2 else GYOKER / "data" / "adatok.js"

FELTETELEZES = re.compile(r"feltételez", re.IGNORECASE)


def szoveg(v):
    return "" if v is None else str(v).strip()


def szam(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def kerekit(v, jegy=4):
    return None if v is None else round(v, jegy)


def tarhelyek(ws):
    fejlec = [szoveg(c.value) for c in ws[1]]
    oszlop = {nev: i for i, nev in enumerate(fejlec)}

    def ertek(sor, nev):
        return sor[oszlop[nev]] if nev in oszlop else None

    helyek = []
    jeloletlen_szamlalo = {}
    for sorszam, sor in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None for v in sor):
            continue
        kod = szoveg(ertek(sor, "Helykód"))
        zona = szoveg(ertek(sor, "Zóna"))
        sor_nev = szoveg(ertek(sor, "Sor"))
        hely = ertek(sor, "Hely")
        jeloletlen = (not kod) or kod.startswith("(")
        if jeloletlen:
            # A táblázatban nincs kód: ideiglenes azonosítót képzünk (lásd README).
            elotag = "ALSO" if "alsó" in sor_nev else "ROVID" if "rövid" in sor_nev else "JEL"
            kulcs = (zona, elotag)
            jeloletlen_szamlalo[kulcs] = jeloletlen_szamlalo.get(kulcs, 0) + 1
            sorsz = int(hely) if isinstance(hely, (int, float)) else jeloletlen_szamlalo[kulcs]
            kod = f"{zona}-{elotag}-{sorsz:02d}"
        helyek.append({
            "kod": kod,
            "eredetiKod": szoveg(ertek(sor, "Helykód")),
            "jeloletlen": jeloletlen,
            "zona": zona,
            "sor": sor_nev,
            "hely": hely,
            "rajzon": szoveg(ertek(sor, "Rajzon")),
            "x": kerekit(szam(ertek(sor, "X (m)"))),
            "y": kerekit(szam(ertek(sor, "Y (m)"))),
            "folyoso": szoveg(ertek(sor, "Folyosó")),
            "szarny": szoveg(ertek(sor, "Szárny")),
            "pozicio": kerekit(szam(ertek(sor, "Pozíció a folyosón (m)"))),
            "folyosoveg": kerekit(szam(ertek(sor, "Folyosóvég (m)"))),
            "gerincPont": kerekit(szam(ertek(sor, "Gerinc-pont (m)"))),
            "tav": kerekit(szam(ertek(sor, "Táv a kiindulóponttól (m)")), 2),
            "excelSor": sorszam,
        })
    return helyek


def parameterek(ws):
    lista = []
    for sorszam, sor in enumerate(ws.iter_rows(values_only=True), start=1):
        nev = szoveg(sor[0])
        if not nev or sorszam <= 3 and sor[1] is None:
            continue
        if nev in ("Paraméter", "Folyosós útvonal"):
            continue
        lista.append({
            "nev": nev,
            "ertek": sor[1] if isinstance(sor[1], (int, float)) else szoveg(sor[1]),
            "egyseg": szoveg(sor[2]) if len(sor) > 2 else "",
            "megjegyzes": szoveg(sor[3]) if len(sor) > 3 else "",
            "cella": f"B{sorszam}",
            "excelSor": sorszam,
        })
    return lista


def folyosok_lap(ws):
    sorok = list(ws.iter_rows(values_only=True))
    leiras = szoveg(sorok[1][0]) if len(sorok) > 1 else ""
    szarnyak, folyosok, hozzarendeles, feltetelezesek = [], [], [], []
    mod = None
    lista_mod = False
    for sorszam, sor in enumerate(sorok, start=1):
        a = szoveg(sor[0])
        if a == "Szárny":
            mod = "szarny"
            continue
        if a == "Folyosó":
            mod = "folyoso"
            # a J–K oszlop (sor → folyosó hozzárendelés) ettől a sortól indul
        if len(sor) > 10 and sor[9] is not None and szoveg(sor[9]) != "Zóna|Sor (Tárhelylista B|C)":
            hozzarendeles.append({"zonaSor": szoveg(sor[9]), "folyoso": szoveg(sor[10])})
        if a == "Folyosó":
            continue
        if a.startswith("Feltételezések"):
            lista_mod = True
            continue
        if lista_mod:
            if a.startswith("•"):
                feltetelezesek.append({"szoveg": a.lstrip("•").strip(), "excelSor": sorszam})
            continue
        if not a:
            continue
        if mod == "szarny":
            szarnyak.append({
                "szarny": a,
                "gerincPont": szam(sor[1]),
                "kapuTav": kerekit(szam(sor[2])),
                "megjegyzes": szoveg(sor[3]),
            })
        elif mod == "folyoso":
            folyosok.append({
                "id": a,
                "szarny": szoveg(sor[1]),
                "tipus": szoveg(sor[2]),
                "tengely": szoveg(sor[3]),
                "c": kerekit(szam(sor[4])),
                "e": kerekit(szam(sor[5])),
                "kozvetlen": szoveg(sor[6]),
                "megjegyzes": szoveg(sor[7]),
                "excelSor": sorszam,
            })
    return leiras, szarnyak, folyosok, hozzarendeles, feltetelezesek


def feltetelezes_cellak(wb):
    """Minden cella, amelynek a szövegében szerepel a 'feltételez' szótő."""
    talalat = []
    for ws in wb.worksheets:
        for sor in ws.iter_rows():
            for c in sor:
                if isinstance(c.value, str) and FELTETELEZES.search(c.value):
                    talalat.append({"lap": ws.title, "cella": c.coordinate, "szoveg": c.value.strip()})
                if c.comment and FELTETELEZES.search(c.comment.text or ""):
                    talalat.append({"lap": ws.title, "cella": c.coordinate, "szoveg": c.comment.text.strip(),
                                    "cellaMegjegyzes": True})
    return talalat


def main():
    wb = openpyxl.load_workbook(BEMENET, data_only=True)
    helyek = tarhelyek(wb["Tárhelyek"])
    params = parameterek(wb["Paraméterek"])
    leiras, szarnyak, folyosok, hozzarendeles, feltetelezesek = folyosok_lap(wb["Folyosók"])

    def param(kezdet):
        for p in params:
            if p["nev"].startswith(kezdet):
                return p["ertek"]
        raise SystemExit(f"Hiányzó paraméter: {kezdet}")

    kulcsok = {
        "kiindulo": {"x": param("Kiindulópont X"), "y": param("Kiindulópont Y")},
        "atjaro": {"x": param("Átjáró a felső szárnyba X"), "y": param("Átjáró a felső szárnyba Y")},
        "t1KeresztX": param("T1 keresztfolyosó X"),
        "t1FejY": param("T1 fejfolyosó Y"),
        "alsoSavY": param("Felső szárny alsó sáv Y"),
    }
    # A gerinceknél a 'Folyosó helye' oszlopban nincs szám: a Paraméterek lapról pótoljuk.
    for f in folyosok:
        if f["c"] is None and f["tipus"] == "gerinc":
            f["c"] = kulcsok["t1KeresztX"] if f["szarny"] == "T1" else kulcsok["alsoSavY"]
            f["cForras"] = "Paraméterek"

    kodok = [h["kod"] for h in helyek]
    if len(set(kodok)) != len(kodok):
        dupla = sorted({k for k in kodok if kodok.count(k) > 1})
        raise SystemExit(f"Ismétlődő helykódok: {dupla}")

    adat = {
        "forras": BEMENET.name,
        "folyosoLeiras": leiras,
        "helyek": helyek,
        "folyosok": folyosok,
        "szarnyak": szarnyak,
        "hozzarendeles": hozzarendeles,
        "parameterek": params,
        "kulcsok": kulcsok,
        "feltetelezesLista": feltetelezesek,
        "feltetelezesCellak": feltetelezes_cellak(wb),
    }
    KIMENET.parent.mkdir(parents=True, exist_ok=True)
    with open(KIMENET, "w", encoding="utf-8") as f:
        f.write("// AUTOMATIKUSAN GENERÁLT FÁJL – ne szerkeszd kézzel!\n")
        f.write(f"// Forrás: {BEMENET.name}  |  Újragenerálás: python3 tools/xlsx_to_adatok.py\n")
        f.write("window.RAKTAR_ADATOK = ")
        json.dump(adat, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    print(f"{len(helyek)} tárhely, {len(folyosok)} folyosó, "
          f"{len(adat['feltetelezesCellak'])} feltételezés-cella -> {KIMENET}")


if __name__ == "__main__":
    main()
